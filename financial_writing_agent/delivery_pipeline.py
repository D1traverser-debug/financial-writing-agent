from __future__ import annotations

import hashlib
import hmac
import json
import re
import secrets
from pathlib import Path
from typing import Any, Mapping

from docx import Document

from .checks import check_docx, sha256_file
from .delivery_contract import CANONICAL_DELIVERY_PATH, FORBIDDEN_VISIBLE_CHARACTERS, TEXT_SUBSTITUTIONS
from .document import (
    ArticleDocument,
    AuditLine,
    Block,
    build_article_docx,
    build_audit_docx,
    build_interaction_pack_docx,
)
from .models import ArtifactKind
from .private_research import validate_private_research_brief
from .resources import load_packaged_policy


DELIVERY_EVIDENCE_SCHEMA_VERSION = "1.1"
BUILDER_RECEIPT_VERSION = "1.0"
# Process-local issuer key: visible delivery is intentionally current-session only.
# This is provenance/anti-bypass assurance, not a cryptographic trust boundary against
# a malicious caller with arbitrary in-process Python introspection.
_SESSION_ISSUER_KEY = secrets.token_bytes(32)


class CanonicalDeliveryError(ValueError):
    pass


def _require(condition: bool, message: str) -> None:
    if not condition:
        raise CanonicalDeliveryError(message)


def _canonical_json(value: Mapping[str, Any]) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


def _canonical_digest(value: Mapping[str, Any]) -> str:
    return hashlib.sha256(_canonical_json(value)).hexdigest()


def _issuer_sign(value: Mapping[str, Any]) -> str:
    return hmac.new(_SESSION_ISSUER_KEY, _canonical_json(value), hashlib.sha256).hexdigest()


def _issue_build_receipt(payload: Mapping[str, Any]) -> dict[str, Any]:
    body = {"builder_receipt_version": BUILDER_RECEIPT_VERSION, **dict(payload)}
    return {**body, "builder_attestation": _issuer_sign(body)}


def _verify_build_receipt(receipt: Mapping[str, Any]) -> dict[str, Any]:
    _require(receipt.get("builder_receipt_version") == BUILDER_RECEIPT_VERSION, "build receipt version is not canonical")
    signature = str(receipt.get("builder_attestation", ""))
    _require(len(signature) == 64, "build receipt lacks builder attestation")
    body = {key: value for key, value in receipt.items() if key != "builder_attestation"}
    _require(hmac.compare_digest(signature, _issuer_sign(body)), "build receipt was not issued by the current canonical builder session")
    return body


def _word_policy(policy: Mapping[str, Any]) -> Mapping[str, Any]:
    word = policy.get("word", {})
    return word if isinstance(word, Mapping) else {}


def _text_substitutions(policy: Mapping[str, Any]) -> dict[str, str]:
    merged = dict(TEXT_SUBSTITUTIONS)
    configured = _word_policy(policy).get("text_substitutions", {})
    if isinstance(configured, Mapping):
        merged.update({str(source): str(target) for source, target in configured.items() if str(source)})
    return merged


def _forbidden_characters(policy: Mapping[str, Any]) -> tuple[str, ...]:
    values = list(FORBIDDEN_VISIBLE_CHARACTERS)
    configured = _word_policy(policy).get("forbidden_characters", [])
    if isinstance(configured, list):
        values.extend(str(item) for item in configured if str(item))
    return tuple(dict.fromkeys(values))


def normalize_delivery_text(text: str, policy: Mapping[str, Any]) -> str:
    """Apply deterministic user-owned presentation substitutions before Word construction."""
    normalized = str(text)
    for source, target in _text_substitutions(policy).items():
        normalized = normalized.replace(source, target)
    return normalized


def normalize_article_for_delivery(article: ArticleDocument, policy: Mapping[str, Any]) -> ArticleDocument:
    # Company names must be normalized by the same visible-text transform as article
    # prose; otherwise auto-bold/style checks can silently search for stale names.
    return ArticleDocument(
        title=normalize_delivery_text(article.title, policy),
        stock_names=[normalize_delivery_text(name, policy) for name in article.stock_names],
        blocks=[Block(kind=block.kind, text=normalize_delivery_text(block.text, policy)) for block in article.blocks],
    )


def _docx_visible_text(path: str | Path) -> str:
    doc = Document(path)
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _forbidden_character_errors(path: str | Path, policy: Mapping[str, Any]) -> list[str]:
    text = _docx_visible_text(path)
    return [f"forbidden delivery character present: {token!r}" for token in _forbidden_characters(policy) if token in text]


def _semantic_anchor_errors(
    article: ArticleDocument,
    required_bold_anchors: list[str],
    policy: Mapping[str, Any],
    *,
    artifact_kind: ArtifactKind,
) -> list[str]:
    if artifact_kind not in {ArtifactKind.FINANCE_DEEP_DOCX, ArtifactKind.FINANCE_BROAD_DOCX}:
        return []
    style = policy.get("style", {})
    if not isinstance(style, Mapping) or not style.get("bold_skim_path_required", False):
        return []
    errors: list[str] = []
    anchors = [str(anchor).strip() for anchor in required_bold_anchors if str(anchor).strip()]
    if not anchors:
        errors.append("public long-form delivery requires explicit semantic bold anchors")
        return errors
    visible = "\n".join([article.title, *[block.text for block in article.blocks]])
    numeric_material = bool(re.search(r"\d(?:[\d,.]*)(?:%|％|亿元|万元|万亿元|美元|欧元|元|家|个|倍|年|月|日)", visible))
    if numeric_material and not any(re.search(r"\d", anchor) for anchor in anchors):
        errors.append("public long-form with material numeric facts requires at least one numeric bold anchor")
    return errors


def _generic_build_receipt(
    path: str | Path,
    *,
    artifact_kind: ArtifactKind,
    policy: Mapping[str, Any],
    source_quality_evidence_sha256: str | None = None,
) -> dict[str, Any]:
    # Opening the document is a minimal structural smoke check for non-article output
    # types; product-specific semantic QA remains upstream of Delivery.
    _docx_visible_text(path)
    errors = _forbidden_character_errors(path, policy)
    _require(not errors, f"canonical delivery build failed: {errors}")
    payload: dict[str, Any] = {
        "schema_version": DELIVERY_EVIDENCE_SCHEMA_VERSION,
        "canonical_delivery_path": CANONICAL_DELIVERY_PATH,
        "artifact_kind": artifact_kind.value,
        "file_name": Path(path).name,
        "sha256": sha256_file(path),
        "mechanical_qa": "PASS",
        "constraint_qa": "PASS",
    }
    if source_quality_evidence_sha256 is not None:
        payload["source_quality_evidence_sha256"] = source_quality_evidence_sha256
    return _issue_build_receipt(payload)


def build_verified_article_docx(
    article: ArticleDocument,
    output: str | Path,
    *,
    artifact_kind: ArtifactKind,
    required_bold_anchors: list[str],
    policy: dict[str, Any] | None = None,
    require_risk_disclaimer: bool = True,
    source_quality_evidence_sha256: str | None = None,
) -> dict[str, Any]:
    """Canonical owner-native path for a user-visible long-form/private Word build."""
    active_policy = policy or load_packaged_policy()
    normalized = normalize_article_for_delivery(article, active_policy)
    anchor_errors = _semantic_anchor_errors(normalized, required_bold_anchors, active_policy, artifact_kind=artifact_kind)
    _require(not anchor_errors, f"delivery semantic-anchor admission failed: {anchor_errors}")

    path = build_article_docx(normalized, output, active_policy)
    report = check_docx(
        path,
        stock_names=list(normalized.stock_names),
        policy=active_policy,
        required_bold_anchors=list(required_bold_anchors),
        require_deep_navigation=artifact_kind == ArtifactKind.FINANCE_DEEP_DOCX,
        require_risk_disclaimer=require_risk_disclaimer,
    )
    errors = list(report.get("errors", [])) + _forbidden_character_errors(path, active_policy)
    _require(not errors, f"canonical delivery build failed: {errors}")

    payload: dict[str, Any] = {
        "schema_version": DELIVERY_EVIDENCE_SCHEMA_VERSION,
        "canonical_delivery_path": CANONICAL_DELIVERY_PATH,
        "artifact_kind": artifact_kind.value,
        "file_name": Path(path).name,
        "sha256": sha256_file(path),
        "mechanical_qa": "PASS",
        "constraint_qa": "PASS",
        "required_bold_anchor_count": len(required_bold_anchors),
        "paragraphs": report.get("paragraphs"),
        "tables": report.get("tables"),
    }
    if source_quality_evidence_sha256 is not None:
        payload["source_quality_evidence_sha256"] = source_quality_evidence_sha256
    return _issue_build_receipt(payload)


def build_verified_private_research_docx(
    *,
    title: str,
    blocks: list[Block],
    stock_names: list[str],
    output: str | Path,
    admitted_private_research_brief: Mapping[str, Any],
    required_bold_anchors: list[str] | None = None,
    policy: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Render only an admitted PrivateResearchBrief through canonical Delivery."""
    admitted = validate_private_research_brief(admitted_private_research_brief)
    article = ArticleDocument(title=title, blocks=blocks, stock_names=stock_names)
    return build_verified_article_docx(
        article,
        output,
        artifact_kind=ArtifactKind.PRIVATE_RESEARCH_DOCX,
        required_bold_anchors=list(required_bold_anchors or []),
        policy=policy,
        require_risk_disclaimer=False,
        source_quality_evidence_sha256=hashlib.sha256(
            json.dumps(admitted, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
        ).hexdigest(),
    )


def build_verified_audit_docx(
    *,
    title: str,
    lines: list[AuditLine],
    artifacts: list[str],
    output: str | Path,
    summary_lines: list[str] | None = None,
    policy: dict[str, Any] | None = None,
) -> dict[str, Any]:
    active_policy = policy or load_packaged_policy()
    path = build_audit_docx(title, lines, artifacts, output, summary_lines=summary_lines)
    return _generic_build_receipt(path, artifact_kind=ArtifactKind.AUDIT_DOCX, policy=active_policy)


def build_verified_interaction_pack_docx(
    pack: Any,
    output: str | Path,
    *,
    policy: dict[str, Any] | None = None,
) -> dict[str, Any]:
    active_policy = policy or load_packaged_policy()
    path = build_interaction_pack_docx(pack, output, active_policy)
    return _generic_build_receipt(path, artifact_kind=ArtifactKind.PLATFORM_INTERACTION_PACK_DOCX, policy=active_policy)


def finalize_delivery_evidence(
    build_receipt: Mapping[str, Any],
    *,
    source_ref: str,
    source_identity: str,
    render_qa: str,
    visual_qa: str,
) -> dict[str, Any]:
    """Bind an issuer-attested canonical build to current render/visual evidence."""
    unsigned_receipt = _verify_build_receipt(build_receipt)
    _require(build_receipt.get("canonical_delivery_path") == CANONICAL_DELIVERY_PATH, "build receipt is not canonical")
    _require(build_receipt.get("mechanical_qa") == "PASS", "mechanical_qa must PASS")
    _require(build_receipt.get("constraint_qa") == "PASS", "constraint_qa must PASS")
    _require(render_qa == "PASS", "render_qa must PASS")
    _require(visual_qa == "PASS", "visual_qa must PASS")
    _require(bool(str(source_ref).strip()), "source_ref is required")
    _require(bool(str(source_identity).strip()), "source_identity is required")

    envelope: dict[str, Any] = {
        "schema_version": DELIVERY_EVIDENCE_SCHEMA_VERSION,
        "canonical_delivery_path": CANONICAL_DELIVERY_PATH,
        "artifact_kind": str(build_receipt.get("artifact_kind")),
        "artifact_sha256": str(build_receipt.get("sha256", "")).lower(),
        "build_receipt_sha256": _canonical_digest(unsigned_receipt),
        "builder_attestation": str(build_receipt.get("builder_attestation")),
        "mechanical_qa": "PASS",
        "constraint_qa": "PASS",
        "render_qa": "PASS",
        "visual_qa": "PASS",
        "source_ref": str(source_ref),
        "source_identity": str(source_identity),
    }
    if build_receipt.get("source_quality_evidence_sha256") is not None:
        envelope["source_quality_evidence_sha256"] = str(build_receipt["source_quality_evidence_sha256"])
    signed = {**envelope, "session_attestation": _issuer_sign(envelope)}
    return {"delivery_evidence": signed, "delivery_evidence_sha256": _canonical_digest(signed)}


def validate_delivery_evidence(resource: Mapping[str, Any]) -> None:
    evidence = resource.get("delivery_evidence")
    _require(isinstance(evidence, Mapping), "delivery_evidence is required")
    expected_digest = resource.get("delivery_evidence_sha256")
    _require(isinstance(expected_digest, str) and len(expected_digest) == 64, "delivery_evidence_sha256 is required")
    _require(_canonical_digest(evidence) == expected_digest.lower(), "delivery_evidence digest mismatch")
    session_attestation = str(evidence.get("session_attestation", ""))
    _require(len(session_attestation) == 64, "delivery evidence lacks current-session attestation")
    unsigned = {key: value for key, value in evidence.items() if key != "session_attestation"}
    _require(hmac.compare_digest(session_attestation, _issuer_sign(unsigned)), "delivery evidence was not finalized by the current canonical Delivery session")
    _require(evidence.get("canonical_delivery_path") == CANONICAL_DELIVERY_PATH, "visible artifact bypassed canonical Delivery")
    _require(evidence.get("artifact_kind") == resource.get("kind"), "delivery evidence artifact kind mismatch")
    _require(str(evidence.get("artifact_sha256", "")).lower() == str(resource.get("sha256", "")).lower(), "delivery evidence artifact hash mismatch")
    for gate in ("mechanical_qa", "constraint_qa", "render_qa", "visual_qa"):
        _require(evidence.get(gate) == "PASS", f"delivery evidence requires {gate}=PASS")
    _require(bool(str(evidence.get("source_ref", "")).strip()), "delivery evidence source_ref is required")
    _require(bool(str(evidence.get("source_identity", "")).strip()), "delivery evidence source_identity is required")
