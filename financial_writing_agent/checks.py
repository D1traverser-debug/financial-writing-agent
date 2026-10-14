from __future__ import annotations

import hashlib
import json
import math
import re
from pathlib import Path
from typing import Any, Iterable

from docx import Document

from .models import ArtifactKind, TaskContract


_TEMPLATE_SIGNAL_PATTERNS: dict[str, re.Pattern[str]] = {
    "not_x_but_y": re.compile(r"(?:真正(?:的)?[^。！？\n]{0,24})?(?:不是|不在)[^。！？\n]{0,32}(?:而是|而在)"),
    "if_only_look": re.compile(r"如果只看"),
    "this_is_exactly": re.compile(r"这正是"),
    "most_worth_staging": re.compile(r"最值得(?:关注|写清楚|看的|讨论|留意)"),
    "core_is": re.compile(r"核心(?:在于|是)"),
    "worth_noting": re.compile(r"值得注意的是"),
    "in_other_words": re.compile(r"换句话说"),
    "problem_is": re.compile(r"问题在于"),
    "surface_real_change": re.compile(r"真正(?:的)?变化"),
    "surface_not_simple": re.compile(r"不是[^。！？\n]{0,24}(?:这么|那么)?简单"),
}

# Publication copy should not reveal the research trail merely to signal diligence.
# Research/audit artifacts retain exact provenance. Subject/actor attribution remains
# a semantic judgment, but source-framing language below is a deterministic hard gate.
_PUBLICATION_SOURCE_PATTERNS: dict[str, re.Pattern[str]] = {
    "research_report_attribution": re.compile(r"(?:研报|研究报告)[^。！？\n]{0,18}(?:指出|认为|显示|测算|预计|提到|称)"),
    "broker_research_attribution": re.compile(r"(?:证券|券商|研究所|机构)[^。！？\n]{0,18}(?:研报|研究报告)[^。！？\n]{0,18}(?:指出|认为|显示|测算|预计|提到|称)?"),
    "according_to_report": re.compile(r"据[^，。！？\n]{1,30}(?:报道|消息|资料|数据|研报|研究报告)"),
    "public_material_framing": re.compile(r"公开(?:资料|报道|信息|数据)[^。！？\n]{0,12}(?:显示|指出|提到|称|梳理)?"),
    "material_shows": re.compile(r"(?:资料|报道)显示"),
}

_GENERIC_DECISION_CONDITIONAL = re.compile(r"谁(?:能|先|可以)[^。！？\n]{0,28}(?:就|才)[^。！？\n]{0,36}")
_COMPARISON_SIGNAL = re.compile(r"更|高于|低于|快于|慢于|领先|落后|直接|确定|兑现|纯度|弹性|优势|弱于|强于|反例|当前证据")
_MONITOR_HEADING = re.compile(r"验证|兑现|观察|跟踪|后续|下一步|看什么|信号|阶段|判断|检查")
_SOURCE_SECTION_HEADING = re.compile(r"^(?:资料来源|参考资料|参考来源|来源汇总)(?:[：:]|$)")


def sha256_file(path: str | Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def expected_artifact_kinds(contract: TaskContract) -> set[ArtifactKind]:
    return set(contract.required_outputs)


def check_required_outputs(contract: TaskContract, artifacts: dict[ArtifactKind, str]) -> list[str]:
    errors: list[str] = []
    for kind in expected_artifact_kinds(contract):
        path = artifacts.get(kind)
        if not path:
            errors.append(f"missing artifact mapping: {kind}")
        elif not Path(path).is_file():
            errors.append(f"artifact does not exist: {kind} -> {path}")
    return errors


def scan_finance_redlines(text: str, phrases: Iterable[str]) -> list[str]:
    return sorted({p for p in phrases if p in text})


def scan_publication_source_leaks(text: str) -> list[str]:
    """Return research-provenance/source-framing families present in publication copy."""
    return sorted(name for name, pattern in _PUBLICATION_SOURCE_PATTERNS.items() if pattern.search(text))


def analyze_template_scaffolding(text: str) -> dict[str, Any]:
    """Detect clustered rhetorical scaffolding without treating normal words as banned tokens.

    One occurrence is not a defect. The guard trips only when the same stage-like construction
    recurs or several families cluster densely enough to make the prose mechanically patterned.
    """
    counts = {name: len(pattern.findall(text)) for name, pattern in _TEMPLATE_SIGNAL_PATTERNS.items()}
    active = {name: count for name, count in counts.items() if count}
    total = sum(active.values())
    chars = max(len(text), 1)
    density_per_1000 = round(total * 1000 / chars, 2)
    repeated_family = any(count >= 3 for count in active.values())
    clustered = total >= 5 and len(active) >= 3 and density_per_1000 >= 1.5
    return {
        "counts": active,
        "total": total,
        "families": len(active),
        "density_per_1000_chars": density_per_1000,
        "requires_revise": repeated_family or clustered,
    }


def analyze_publication_scaffolding(payload: Any) -> dict[str, Any]:
    """Run template analysis across title, H1s and body, not body prose alone."""
    title = str(getattr(payload, "title", ""))
    blocks = list(getattr(payload, "blocks", []))
    h1s = [str(getattr(block, "text", "")) for block in blocks if getattr(block, "kind", None) == "h1"]
    body = "\n".join(str(getattr(block, "text", "")) for block in blocks)
    combined = "\n".join([title, *h1s, body])
    result = analyze_template_scaffolding(combined)
    result["surface_title"] = title
    result["surface_h1_count"] = len(h1s)
    return result


def payload_text(payload: Any) -> str:
    return "\n".join(getattr(block, "text", "") for block in getattr(payload, "blocks", []))


def check_payload_navigation(payload: Any, *, deep: bool) -> list[str]:
    """Protect real navigation, including an image-safe section-entry bridge.

    Every H1 that is followed by lower-level headings must first have a substantive lead/body
    paragraph. That paragraph explains why the following branches matter and also gives the user
    a safe place to insert an H1-related image without placing it directly against H1/H2 styles.
    Deep additionally rejects extreme monolithic body blocks.
    """
    errors: list[str] = []
    blocks = list(getattr(payload, "blocks", []))
    for index, block in enumerate(blocks):
        kind = getattr(block, "kind", None)
        if kind == "h1" and index + 1 < len(blocks):
            next_kind = getattr(blocks[index + 1], "kind", None)
            if next_kind in {"h2", "h3"}:
                errors.append(f"H1 must have a substantive image-safe lead paragraph before {next_kind.upper()}: block={index}")
            elif next_kind in {"lead", "body"}:
                lead_text = re.sub(r"\*\*", "", str(getattr(blocks[index + 1], "text", ""))).strip()
                if len(lead_text) < 40 and index + 2 < len(blocks) and getattr(blocks[index + 2], "kind", None) in {"h2", "h3"}:
                    errors.append(f"H1 lead paragraph is too thin to bridge into lower-level headings: block={index}, chars={len(lead_text)}")

        if deep and kind == "body":
            text = str(getattr(block, "text", ""))
            if len(text) >= 1400:
                errors.append(f"FINANCE_DEEP monolithic body block requires structural rewrite: block={index}, chars={len(text)}")
    return errors


def check_company_decision_utility(payload: Any, research: Any) -> list[str]:
    """Require a named comparative verdict when at least two MATERIAL companies exist."""
    if payload is None or research is None:
        return []
    companies = list(getattr(research, "universe", []))
    material = [company for company in companies if getattr(company, "materiality", None) == "MATERIAL"]
    if len(material) < 2:
        return []

    errors: list[str] = []
    decision_map = getattr(research, "company_decision_map", None)
    if decision_map is None:
        errors.append("multi-company MATERIAL research requires company_decision_map")

    material_names = [company.name for company in material]
    blocks = list(getattr(payload, "blocks", []))
    paragraphs = [str(getattr(block, "text", "")) for block in blocks if getattr(block, "kind", None) in {"lead", "body"}]
    comparative = [
        paragraph
        for paragraph in paragraphs
        if sum(name in paragraph for name in material_names) >= 2 and _COMPARISON_SIGNAL.search(paragraph)
    ]
    if not comparative:
        errors.append("multi-company article lacks a named evidence-bounded comparative verdict paragraph")

    for paragraph in paragraphs:
        if _GENERIC_DECISION_CONDITIONAL.search(paragraph) and not any(name in paragraph for name in material_names):
            errors.append("generic zero-information company-selection conditional must be replaced by named evidence-bounded comparison")
            break
    return errors


def check_deep_company_grounding(payload: Any, research: Any) -> list[str]:
    """Cross-check the material research universe against the FINANCE_DEEP draft."""
    if payload is None or research is None:
        return []
    text = payload_text(payload)
    companies = list(getattr(research, "universe", []))
    material = [company for company in companies if getattr(company, "materiality", None) == "MATERIAL"]
    missing = [company.name for company in material if company.name not in text]
    errors = [f"FINANCE_DEEP silently omitted MATERIAL company: {name}" for name in missing]

    sections: list[tuple[str, str]] = []
    current_heading = ""
    current_parts: list[str] = []
    for block in getattr(payload, "blocks", []):
        kind = getattr(block, "kind", None)
        block_text = getattr(block, "text", "")
        if kind == "h1":
            if current_heading:
                sections.append((current_heading, "\n".join(current_parts)))
            current_heading = block_text
            current_parts = []
        elif current_heading:
            current_parts.append(block_text)
    if current_heading:
        sections.append((current_heading, "\n".join(current_parts)))

    material_names = [company.name for company in material]
    for heading, body in sections:
        if _MONITOR_HEADING.search(heading) and material_names:
            section_text = f"{heading}\n{body}"
            if not any(name in section_text for name in material_names):
                errors.append(f"monitoring/validation section is detached from the stock universe: {heading}")
    errors.extend(check_company_decision_utility(payload, research))
    return errors


def _is_bold_occurrence(doc: Document, needle: str) -> bool:
    if not needle:
        return False
    for paragraph in doc.paragraphs:
        full = "".join(run.text for run in paragraph.runs)
        if needle not in full:
            continue
        mask: list[bool] = []
        for run in paragraph.runs:
            mask.extend([bool(run.bold)] * len(run.text))
        start = 0
        while True:
            idx = full.find(needle, start)
            if idx < 0:
                break
            if all(mask[idx : idx + len(needle)]):
                return True
            start = idx + 1
    return False


def _unbold_occurrence_count(doc: Document, needle: str) -> int:
    if not needle:
        return 0
    unbold = 0
    for paragraph in doc.paragraphs:
        full = "".join(run.text for run in paragraph.runs)
        if needle not in full:
            continue
        mask: list[bool] = []
        for run in paragraph.runs:
            mask.extend([bool(run.bold)] * len(run.text))
        start = 0
        while True:
            idx = full.find(needle, start)
            if idx < 0:
                break
            if not all(mask[idx : idx + len(needle)]):
                unbold += 1
            start = idx + 1
    return unbold


def _check_run_sizes(paragraph, expected_pt: float, label: str, errors: list[str]) -> None:
    visible_runs = [run for run in paragraph.runs if run.text]
    if not visible_runs:
        return
    for run in visible_runs:
        size = run.font.size
        if size is None or not math.isclose(size.pt, expected_pt, rel_tol=0.0, abs_tol=0.01):
            actual = None if size is None else round(size.pt, 2)
            errors.append(f"{label} typography drift: expected {expected_pt}pt, got {actual}pt in {paragraph.text[:60]!r}")
            return


def _check_typography(doc: Document, policy: dict, errors: list[str]) -> None:
    word = policy.get("word", {})
    title_pt = float(word["title_pt"])
    h1_pt = float(word["h1_pt"])
    h2_pt = float(word["h2_pt"])
    h3_pt = float(word["h3_pt"])
    body_pt = float(word["body_pt"])

    nonempty = [p for p in doc.paragraphs if p.text.strip()]
    if nonempty:
        _check_run_sizes(nonempty[0], title_pt, "title", errors)

    circled = re.compile(r"^[①-⑳]")
    h2_prefix = re.compile(r"^（\d+）")
    for p in nonempty[1:]:
        if p.style and p.style.name == "Heading 1":
            _check_run_sizes(p, h1_pt, "H1", errors)
        elif h2_prefix.match(p.text.strip()):
            _check_run_sizes(p, h2_pt, "H2", errors)
        elif circled.match(p.text.strip()):
            _check_run_sizes(p, h3_pt, "H3", errors)
        else:
            _check_run_sizes(p, body_pt, "body", errors)


def check_docx(
    path: str | Path,
    *,
    stock_names: list[str],
    policy: dict,
    required_bold_anchors: list[str] | None = None,
    require_deep_navigation: bool = False,
    require_risk_disclaimer: bool = True,
) -> dict:
    """Run deterministic Word checks without treating formatting density as editorial quality."""
    _ = require_deep_navigation
    doc = Document(path)
    errors: list[str] = []
    warnings: list[str] = []

    heading2_plus = [p.text for p in doc.paragraphs if p.style and p.style.name.startswith("Heading ") and p.style.name != "Heading 1"]
    if heading2_plus:
        errors.append(f"true Heading 2+ styles found: {heading2_plus[:5]}")

    h1s = [p.text.strip() for p in doc.paragraphs if p.style and p.style.name == "Heading 1"]
    for text in h1s:
        if text.startswith(("·", "•")):
            errors.append(f"forbidden heading hierarchy prefix: {text[:40]}")
        if not re.match(r"^[一二三四五六七八九十百]+、", text):
            errors.append(f"H1 missing Chinese sequence prefix: {text[:60]}")
        if "风险" in text and ("免责声明" in text or "提示" in text):
            errors.append("risk disclaimer must be an inline italic paragraph, not a numbered H1")

    lower_numbered = [p.text.strip() for p in doc.paragraphs if re.match(r"^(?:（\d+）|[①-⑳])", p.text.strip())]

    _check_typography(doc, policy, errors)

    if policy.get("word", {}).get("auto_bold_company_names", False):
        for stock_name in sorted({name for name in stock_names if name}):
            unbold = _unbold_occurrence_count(doc, stock_name)
            if unbold:
                errors.append(f"A-share/company name must be bold at every occurrence: {stock_name} unbold_occurrences={unbold}")

    if not policy.get("word", {}).get("allow_underline", False):
        underlined = [p.text[:60] for p in doc.paragraphs if any(run.text and bool(run.underline) for run in p.runs)]
        if underlined:
            errors.append(f"publication article contains prohibited underline(s): {underlined[:5]}")

    if not policy.get("word", {}).get("end_source_section_allowed", False):
        source_sections = [p.text.strip() for p in doc.paragraphs if _SOURCE_SECTION_HEADING.match(p.text.strip())]
        if source_sections:
            errors.append(f"publication article must not contain end source section: {source_sections[:3]}")

    text = "\n".join(p.text for p in doc.paragraphs)
    if policy.get("publication_source_privacy", {}).get("required", False):
        source_leaks = scan_publication_source_leaks(text)
        if source_leaks:
            errors.append(f"publication copy exposes research/source provenance: {source_leaks}")

    risk_paragraphs = [p for p in doc.paragraphs if p.text.strip().startswith("风险提示：")]
    if require_risk_disclaimer:
        if not risk_paragraphs:
            errors.append("inline risk disclaimer paragraph not found")
        elif len(risk_paragraphs) > 1:
            errors.append(f"multiple risk disclaimer paragraphs found: {len(risk_paragraphs)}")
        else:
            risk_p = risk_paragraphs[0]
            risk_text = risk_p.text.strip()[len("风险提示：") :].strip()
            if not risk_text:
                errors.append("risk/disclaimer body must not be empty")
            min_chars = int(policy.get("risk_disclaimer", {}).get("min_body_chars", 0) or 0)
            if min_chars and len(risk_text) < min_chars:
                errors.append(f"risk/disclaimer body too thin: {len(risk_text)} < {min_chars} chars")
            recommended_max = int(policy.get("risk_disclaimer", {}).get("recommended_max_chars", 0) or 0)
            if recommended_max and len(risk_text) > recommended_max:
                warnings.append(f"risk/disclaimer body may be verbose: {len(risk_text)} > {recommended_max} chars")
            if policy.get("risk_disclaimer", {}).get("italic", False):
                visible_runs = [run for run in risk_p.runs if run.text]
                if not visible_runs or not all(run.italic is True for run in visible_runs):
                    errors.append("risk disclaimer must be italic")

    if not policy["word"].get("allow_tables", False) and doc.tables:
        errors.append(f"publication article contains prohibited table(s): {len(doc.tables)}")

    for anchor in required_bold_anchors or []:
        if not _is_bold_occurrence(doc, anchor):
            errors.append(f"required skim-path anchor not bold in Word: {anchor}")

    hits = scan_finance_redlines(text, policy["finance_compliance"]["hard_fail_phrases"])
    if hits:
        errors.append(f"actionable-finance redline phrase(s): {hits}")

    extreme_hits = scan_finance_redlines(text, policy.get("style", {}).get("hard_fail_extreme_phrases", []))
    if extreme_hits:
        errors.append(f"unsupported deterministic promotional phrase(s): {extreme_hits}")

    return {
        "ok": not errors,
        "errors": errors,
        "warnings": warnings,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "inline_images": len(doc.inline_shapes),
        "required_bold_anchors": len(required_bold_anchors or []),
        "lower_numbered_headings": len(lower_numbered),
    }


def load_policy(path: str | Path) -> dict:
    candidate = Path(path)
    if candidate.is_file():
        return json.loads(candidate.read_text(encoding="utf-8"))
    if candidate.name == "policy.json":
        from .resources import load_packaged_policy
        return load_packaged_policy()
    raise FileNotFoundError(candidate)
