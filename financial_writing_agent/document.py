from __future__ import annotations

import re
from pathlib import Path
from typing import Literal

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from pydantic import BaseModel, Field


class Block(BaseModel):
    kind: Literal["lead", "h1", "h2", "h3", "body"]
    text: str


class ArticleDocument(BaseModel):
    title: str
    blocks: list[Block]
    stock_names: list[str] = Field(default_factory=list)


def _split_bold(
    text: str,
    *,
    stock_names: list[str] | None = None,
    auto_bold_stock_names: bool = False,
) -> list[tuple[str, bool]]:
    """Parse explicit ``**...**`` emphasis and optional A-share name emphasis."""
    names = sorted({name for name in (stock_names or []) if name}, key=len, reverse=True)
    name_set = set(names)
    name_pattern = None
    if auto_bold_stock_names and names:
        name_pattern = re.compile("(" + "|".join(re.escape(name) for name in names) + ")")

    parts = re.split(r"(\*\*.*?\*\*)", text)
    out: list[tuple[str, bool]] = []
    for part in parts:
        if not part:
            continue
        explicit = part.startswith("**") and part.endswith("**")
        raw = part[2:-2] if explicit else part
        if explicit or name_pattern is None:
            out.append((raw, explicit))
            continue
        for piece in name_pattern.split(raw):
            if not piece:
                continue
            out.append((piece, piece in name_set))
    return out


def _set_cn_font(run, name="Microsoft YaHei"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _format_run(run, *, size_pt: float, bold: bool = False, italic: bool = False, color: str | None = None) -> None:
    run.bold = bold
    run.italic = italic
    run.underline = False
    run.font.size = Pt(size_pt)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    _set_cn_font(run)


def _chinese_number(n: int) -> str:
    digits = "零一二三四五六七八九"
    if n <= 0:
        return str(n)
    if n < 10:
        return digits[n]
    if n == 10:
        return "十"
    if n < 20:
        return "十" + digits[n - 10]
    if n < 100:
        tens, ones = divmod(n, 10)
        return digits[tens] + "十" + (digits[ones] if ones else "")
    return str(n)


def _strip_heading_prefix(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^[·•]\s*", "", text)
    text = re.sub(r"^[一二三四五六七八九十百]+、\s*", "", text)
    text = re.sub(r"^（\d+）\s*", "", text)
    text = re.sub(r"^[①-⑳]\s*", "", text)
    return text.strip()


def _heading_label(kind: str, text: str, h1: int, h2: int, h3: int) -> str:
    clean = _strip_heading_prefix(text)
    if kind == "h1":
        return f"{_chinese_number(h1)}、{clean}"
    if kind == "h2":
        return f"（{h2}）{clean}"
    if kind == "h3":
        circled = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳"
        prefix = circled[h3 - 1] if 1 <= h3 <= len(circled) else f"（{h3}）"
        return f"{prefix}{clean}"
    return clean


def _is_risk_heading(text: str) -> bool:
    clean = _strip_heading_prefix(text)
    return "风险" in clean and ("免责声明" in clean or "提示" in clean)


def build_article_docx(article: ArticleDocument, output: str | Path, policy: dict) -> Path:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.size = Pt(policy["word"]["body_pt"])
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(article.title)
    _format_run(r, bold=True, size_pt=policy["word"]["title_pt"])

    auto_bold_stock_names = bool(policy.get("word", {}).get("auto_bold_company_names", False))

    h1 = h2 = h3 = 0
    risk_mode = False
    for block in article.blocks:
        render_text = block.text

        if block.kind == "h1" and _is_risk_heading(block.text):
            risk_mode = True
            continue
        if block.kind in {"h1", "h2", "h3"} and risk_mode:
            risk_mode = False

        if block.kind == "h1":
            h1 += 1
            h2 = h3 = 0
            render_text = _heading_label("h1", block.text, h1, h2, h3)
            p = doc.add_paragraph(style="Heading 1")
            for text, _ in _split_bold(render_text):
                run = p.add_run(text)
                _format_run(run, bold=True, size_pt=policy["word"]["h1_pt"], color=policy["word"]["h1_rgb"])
        elif block.kind == "h2":
            h2 += 1
            h3 = 0
            render_text = _heading_label("h2", block.text, h1, h2, h3)
            p = doc.add_paragraph(style="Normal")
            for text, _ in _split_bold(render_text):
                run = p.add_run(text)
                _format_run(run, bold=True, size_pt=policy["word"]["h2_pt"])
        elif block.kind == "h3":
            h3 += 1
            render_text = _heading_label("h3", block.text, h1, h2, h3)
            p = doc.add_paragraph(style="Normal")
            for text, _ in _split_bold(render_text):
                run = p.add_run(text)
                _format_run(run, bold=True, size_pt=policy["word"]["h3_pt"])
        else:
            p = doc.add_paragraph(style="Normal")
            is_risk = risk_mode or render_text.strip().startswith("风险提示：")
            if is_risk and not render_text.strip().startswith("风险提示："):
                render_text = f"风险提示：{render_text.strip()}"
            for text, bold in _split_bold(
                render_text,
                stock_names=article.stock_names,
                auto_bold_stock_names=auto_bold_stock_names,
            ):
                run = p.add_run(text)
                _format_run(run, bold=bool(bold), italic=is_risk, size_pt=policy["word"]["body_pt"])

    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


class AuditLine(BaseModel):
    area: str
    status: str
    evidence: str
    notes: str = ""


def build_audit_docx(title: str, lines: list[AuditLine], artifacts: list[str], output: str | Path, summary_lines: list[str] | None = None) -> Path:
    doc = Document()
    h = doc.add_paragraph()
    r = h.add_run(title)
    r.bold = True
    r.underline = False
    r.font.size = Pt(18)
    _set_cn_font(r)

    for summary in summary_lines or []:
        p = doc.add_paragraph(summary)
        for run in p.runs:
            run.underline = False
            _set_cn_font(run)

    for line in lines:
        p = doc.add_paragraph()
        r1 = p.add_run(f"[{line.status}] {line.area}：")
        r1.bold = True
        r1.underline = False
        _set_cn_font(r1)
        r2 = p.add_run(line.evidence)
        r2.underline = False
        _set_cn_font(r2)
        if line.notes:
            r3 = p.add_run(f"｜{line.notes}")
            r3.underline = False
            _set_cn_font(r3)

    p = doc.add_paragraph("最终文件身份：")
    for run in p.runs:
        run.underline = False
    for a in artifacts:
        p = doc.add_paragraph(a, style="List Bullet")
        for run in p.runs:
            run.underline = False

    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


def build_interaction_pack_docx(pack, output: str | Path, policy: dict) -> Path:
    """Render the standalone platform interaction pack as a Word artifact.

    The pack is intentionally not an ArticleDocument: it contains several native short-form
    rewrites rather than one long article and therefore has no article-level risk paragraph or
    heading hierarchy. Finance compliance/source privacy still apply in semantic/mechanical QA.
    """
    doc = Document()
    styles = doc.styles
    body_pt = float(policy["word"]["body_pt"])
    styles["Normal"].font.size = Pt(body_pt)
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"{pack.topic}｜互动短内容包")
    _format_run(run, bold=True, size_pt=policy["word"]["title_pt"])

    auto_bold_stock_names = bool(policy.get("word", {}).get("auto_bold_company_names", False))

    for index, post in enumerate(pack.posts, start=1):
        stock_names = list(getattr(post, "stock_names", []) or [])

        label = doc.add_paragraph(style="Normal")
        run = label.add_run(f"{index}. {post.platform_family}")
        _format_run(run, bold=True, size_pt=body_pt)

        for field_name, field_value, bold_label in (
            ("开头", post.hook, True),
            ("正文", post.body, False),
            ("互动", post.interaction_prompt, True),
        ):
            paragraph = doc.add_paragraph(style="Normal")
            label_run = paragraph.add_run(f"{field_name}：")
            _format_run(label_run, bold=bold_label, size_pt=body_pt)
            for text, bold in _split_bold(
                field_value,
                stock_names=stock_names,
                auto_bold_stock_names=auto_bold_stock_names,
            ):
                text_run = paragraph.add_run(text)
                _format_run(text_run, bold=bool(bold), size_pt=body_pt)

    boundary = doc.add_paragraph(style="Normal")
    boundary_run = boundary.add_run(f"证据边界：{pack.evidence_boundary}")
    _format_run(boundary_run, italic=True, size_pt=body_pt)

    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output
