from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

from financial_writing_agent.document import build_interaction_pack_docx
from financial_writing_agent.resources import load_packaged_policy
from financial_writing_agent.schemas import InteractionPost, PlatformInteractionPack


def _render(root: Path, docx_path: Path, outdir: Path) -> list[Path]:
    proc = subprocess.run(
        [sys.executable, str(root / "scripts" / "render_docx.py"), str(docx_path), "--out", str(outdir)],
        capture_output=True,
        text=True,
    )
    if proc.returncode != 0:
        raise SystemExit(proc.stderr.strip() or proc.stdout.strip() or f"render smoke failed: {docx_path.name}")
    pages = [Path(line.strip()) for line in proc.stdout.splitlines() if line.strip().endswith(".png")]
    if not pages or not all(p.exists() and p.stat().st_size > 0 for p in pages):
        raise SystemExit(f"render smoke did not produce valid PNG pages: {docx_path.name}")
    return pages


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    policy = load_packaged_policy()
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)

        interaction_path = tmp / "interaction_pack_smoke.docx"
        interaction = PlatformInteractionPack(
            topic="Reader Value Render Smoke",
            evidence_boundary="同一已审计研究包，不新增事实。",
            posts=[
                InteractionPost(
                    platform_family="WECHAT_IMAGE_TEXT",
                    hook="同一主题，先看最硬的变化。",
                    body="甲公司当前兑现更直接，乙公司弹性更高但验证链更长。",
                    interaction_prompt="你更看重确定性还是弹性？",
                    stock_names=["甲公司", "乙公司"],
                ),
                InteractionPost(
                    platform_family="TOUTIAO_MICRO",
                    hook="价格变化不是全部。",
                    body="真正影响财务结果的是兑现路径和公司差异。",
                    interaction_prompt="你会先验证收入还是渠道？",
                ),
                InteractionPost(
                    platform_family="BAIJIAHAO_DYNAMIC",
                    hook="同一个概念，逻辑也能分层。",
                    body="当前证据强弱不同，判断应随新数据调整。",
                    interaction_prompt="你最关心哪一个验证信号？",
                ),
            ],
        )
        build_interaction_pack_docx(interaction, interaction_path, policy)
        interaction_pages = _render(root, interaction_path, tmp / "rendered_interaction")

        # Also retain one minimal generic DOCX to prove the rendering pipeline itself is
        # not coupled only to the interaction-pack builder.
        from docx import Document

        generic_path = tmp / "generic_smoke.docx"
        doc = Document()
        doc.add_heading("Financial Writing Agent Render Smoke", level=1)
        doc.add_paragraph("Render validation must produce at least one PNG page.")
        doc.save(generic_path)
        generic_pages = _render(root, generic_path, tmp / "rendered_generic")

        print(
            "PASS: DOCX render pipeline produced "
            f"{len(generic_pages)} generic + {len(interaction_pages)} interaction-pack PNG page(s)"
        )


if __name__ == "__main__":
    main()
