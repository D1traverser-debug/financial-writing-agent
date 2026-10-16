from docx import Document
import pytest

from financial_writing_agent.delivery_pipeline import (
    CanonicalDeliveryError,
    build_verified_article_docx,
    build_verified_private_research_docx,
    finalize_delivery_evidence,
    validate_delivery_evidence,
)
from financial_writing_agent.document import ArticleDocument, Block
from financial_writing_agent.models import ArtifactKind, TaskContract


def _is_all_occurrences_bold(path, needle):
    doc = Document(path)
    found = 0
    for paragraph in doc.paragraphs:
        full = "".join(run.text for run in paragraph.runs)
        if needle not in full:
            continue
        mask = []
        for run in paragraph.runs:
            mask.extend([bool(run.bold)] * len(run.text))
        start = 0
        while True:
            index = full.find(needle, start)
            if index < 0:
                break
            found += 1
            assert all(mask[index:index + len(needle)])
            start = index + 1
    return found


def test_canonical_build_normalizes_forbidden_dot_bolds_company_and_number_and_numbers_headings(tmp_path):
    article = ArticleDocument(
        title="跨境电商·交付回归",
        stock_names=["安克创新"],
        blocks=[
            Block(kind="h1", text="公司差异"),
            Block(kind="body", text="安克创新的收入结构需要和规则一起看，**1.45万亿元**是本段关键数据。"),
            Block(kind="h2", text="兑现路径"),
            Block(kind="body", text="再次观察安克创新。"),
        ],
    )
    output = tmp_path / "deep.docx"
    receipt = build_verified_article_docx(
        article,
        output,
        artifact_kind=ArtifactKind.FINANCE_DEEP_DOCX,
        required_bold_anchors=["1.45万亿元"],
        require_risk_disclaimer=False,
    )

    doc = Document(output)
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "·" not in text
    assert "跨境电商、交付回归" in text
    assert any(paragraph.text.startswith("一、") for paragraph in doc.paragraphs)
    assert any(paragraph.text.startswith("（1）") for paragraph in doc.paragraphs)
    assert _is_all_occurrences_bold(output, "安克创新") == 2
    assert _is_all_occurrences_bold(output, "1.45万亿元") == 1
    assert receipt["canonical_delivery_path"] == "financial_writing_agent.delivery_pipeline"
    assert receipt["mechanical_qa"] == "PASS"
    assert receipt["constraint_qa"] == "PASS"


def test_numeric_longform_cannot_reach_word_without_explicit_numeric_skim_anchor(tmp_path):
    article = ArticleDocument(
        title="测试",
        blocks=[Block(kind="h1", text="数据"), Block(kind="body", text="规模达到1.45万亿元。")],
    )
    with pytest.raises(CanonicalDeliveryError, match="numeric bold anchor"):
        build_verified_article_docx(
            article,
            tmp_path / "bad.docx",
            artifact_kind=ArtifactKind.FINANCE_DEEP_DOCX,
            required_bold_anchors=["规模变化"],
            require_risk_disclaimer=False,
        )


def test_private_research_has_structured_word_artifact_kind_but_does_not_mutate_core_task_outputs(tmp_path):
    contract = TaskContract(prompt="test")
    assert ArtifactKind.PRIVATE_RESEARCH_DOCX not in contract.required_outputs

    output = tmp_path / "private.docx"
    receipt = build_verified_private_research_docx(
        title="私人研究稿",
        stock_names=["安克创新"],
        blocks=[
            Block(kind="h1", text="私人判断"),
            Block(kind="body", text="安克创新·私人研究，仅供内部研究。"),
        ],
        output=output,
    )
    assert receipt["artifact_kind"] == "PRIVATE_RESEARCH_DOCX"
    assert "·" not in "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert _is_all_occurrences_bold(output, "安克创新") == 1


def test_visible_evidence_binds_build_render_visual_and_rejects_tampering():
    build = {
        "canonical_delivery_path": "financial_writing_agent.delivery_pipeline",
        "artifact_kind": "FINANCE_DEEP_DOCX",
        "sha256": "a" * 64,
        "mechanical_qa": "PASS",
        "constraint_qa": "PASS",
    }
    evidence = finalize_delivery_evidence(
        build,
        source_ref="workspaces/articles/x/publication/FINANCE_DEEP.json",
        source_identity="git-blob:source",
        render_qa="PASS",
        visual_qa="PASS",
    )
    resource = {
        "kind": "FINANCE_DEEP_DOCX",
        "sha256": "a" * 64,
        **evidence,
    }
    validate_delivery_evidence(resource)

    tampered = {**resource, "delivery_evidence": {**resource["delivery_evidence"], "visual_qa": "OPEN"}}
    with pytest.raises(CanonicalDeliveryError, match="digest mismatch"):
        validate_delivery_evidence(tampered)
