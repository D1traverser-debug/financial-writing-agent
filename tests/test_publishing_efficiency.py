import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from financial_writing_agent.checks import check_docx
from financial_writing_agent.document import ArticleDocument, Block, build_article_docx


ROOT = Path(__file__).resolve().parents[1]


def _policy():
    return json.loads((ROOT / "config" / "policy.json").read_text(encoding="utf-8"))


def _risk(chars: int = 80) -> str:
    base = "基于公开信息研究讨论，不构成证券投资建议或收益承诺；市场、政策及公司经营均可能变化，结论可能失效，请独立判断并自担风险。"
    if chars <= len(base):
        return base[:chars]
    return base + ("风险" * ((chars - len(base) + 1) // 2))


def _article(body: str, *, risk: str | None = None) -> ArticleDocument:
    return ArticleDocument(
        title="测试标题",
        stock_names=[],
        blocks=[
            Block(kind="lead", text="导读。"),
            Block(kind="h1", text="核心判断"),
            Block(kind="h2", text="实际传导"),
            Block(kind="h3", text="公司证据"),
            Block(kind="body", text=body),
            Block(kind="h1", text="风险提示与免责声明"),
            Block(kind="body", text=risk or _risk()),
        ],
    )


def test_policy_locks_completion_oriented_broad_and_compact_risk_boundary():
    policy = _policy()
    assert policy["editorial"]["finance_deep_density_review_chars"] == 6000
    assert policy["editorial"]["finance_deep_max_chars"] == 10000
    assert policy["editorial"]["finance_broad_preferred_min_chars"] == 900
    assert policy["editorial"]["finance_broad_preferred_max_chars"] == 1500
    assert policy["editorial"]["finance_broad_hard_max_chars"] == 1800
    assert policy["risk_disclaimer"]["min_body_chars"] == 60
    assert policy["risk_disclaimer"]["recommended_max_chars"] == 120
    assert policy["risk_disclaimer"]["hard_max_chars"] == 160
    assert policy["word"]["title_alignment"] == "left"
    assert (ROOT / "config" / "policy.json").read_bytes() == (
        ROOT / "financial_writing_agent" / "resources" / "policy.json"
    ).read_bytes()


def test_word_title_is_left_aligned_and_hierarchy_is_preserved(tmp_path):
    policy = _policy()
    article = _article("公司与股票结论必须落到实际证据。")
    path = build_article_docx(article, tmp_path / "FINANCE_DEEP.docx", policy)
    result = check_docx(path, stock_names=[], policy=policy, expected_images=0)
    assert result["ok"], result

    doc = Document(path)
    assert doc.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert any(p.text.startswith("（1）") for p in doc.paragraphs)
    assert any(p.text.startswith("①") for p in doc.paragraphs)


def test_finance_broad_over_1800_chars_fails_mechanical_qa(tmp_path):
    policy = _policy()
    article = _article("实" * 1900)
    path = build_article_docx(article, tmp_path / "FINANCE_BROAD.docx", policy)
    result = check_docx(path, stock_names=[], policy=policy, expected_images=0)
    assert not result["ok"]
    assert any("FINANCE_BROAD exceeds default completion-oriented hard cap" in x for x in result["errors"])


def test_finance_broad_between_1500_and_1800_warns_but_can_pass(tmp_path):
    policy = _policy()
    article = _article("实" * 1550)
    path = build_article_docx(article, tmp_path / "FINANCE_BROAD.docx", policy)
    result = check_docx(path, stock_names=[], policy=policy, expected_images=0)
    assert result["ok"], result
    assert any("FINANCE_BROAD exceeds preferred distribution range" in x for x in result["warnings"])


def test_finance_deep_over_density_threshold_warns_not_fails(tmp_path):
    policy = _policy()
    article = _article("实" * 6100)
    path = build_article_docx(article, tmp_path / "FINANCE_DEEP.docx", policy)
    result = check_docx(path, stock_names=[], policy=policy, expected_images=0)
    assert result["ok"], result
    assert any("FINANCE_DEEP exceeds density-review threshold" in x for x in result["warnings"])


def test_risk_disclaimer_above_hard_cap_fails(tmp_path):
    policy = _policy()
    article = _article("正文。", risk=_risk(170))
    path = build_article_docx(article, tmp_path / "FINANCE_DEEP.docx", policy)
    result = check_docx(path, stock_names=[], policy=policy, expected_images=0)
    assert not result["ok"]
    assert any("risk/disclaimer body too verbose" in x for x in result["errors"])


def test_editorial_delivery_style_mirrors_lock_new_semantics():
    for name in ("editorial.md", "delivery.md", "style.md"):
        root = (ROOT / "references" / name).read_bytes()
        assert root == (ROOT / "skills" / "financial-writing-agent" / "references" / name).read_bytes()
        assert root == (ROOT / "financial_writing_agent" / "resources" / "references" / name).read_bytes()

    editorial = (ROOT / "references" / "editorial.md").read_text(encoding="utf-8")
    delivery = (ROOT / "references" / "delivery.md").read_text(encoding="utf-8")
    style = (ROOT / "references" / "style.md").read_text(encoding="utf-8")

    assert "shortest complete expression" in editorial
    assert "900–1500 Chinese characters" in editorial
    assert "default hard cap is **1800**" in editorial
    assert "Use and expose this hierarchy proactively" in editorial
    assert "no generic finance filler" in editorial
    assert "left-aligned" in delivery
    assert "60–120 Chinese characters" in delivery
    assert "minimum sufficient disclosure" in delivery
    assert "Anti-genericity / concrete application" in style
    assert "actual company/stock consequences" in style
