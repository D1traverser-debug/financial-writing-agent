import json

from docx import Document

from financial_writing_agent.checks import check_docx
from financial_writing_agent.document import ArticleDocument, Block, build_article_docx


def _risk_text() -> str:
    return (
        "产业需求、订单兑现、价格、竞争格局和公司业绩均可能变化，"
        "文中判断基于公开信息，不构成投资建议。"
    )


def test_word_builder_normalizes_hierarchy_preserves_bold_contract_and_exact_typography(tmp_path):
    policy = json.loads(open("config/policy.json", encoding="utf-8").read())
    article = ArticleDocument(
        title="测试标题",
        stock_names=["国电南瑞"],
        blocks=[
            Block(kind="lead", text="这是导读，涉及国电南瑞，并且**利润开始兑现**。"),
            Block(kind="h1", text="第一部分"),
            Block(kind="h2", text="二级标题"),
            Block(kind="h3", text="三级标题"),
            Block(kind="body", text="国电南瑞在这里再次出现，**订单增长20%**。"),
            Block(kind="h1", text="风险提示与免责声明"),
            Block(kind="body", text=_risk_text()),
        ],
    )
    path = build_article_docx(article, tmp_path / "a.docx", policy)
    result = check_docx(
        path,
        stock_names=article.stock_names,
        policy=policy,
        required_bold_anchors=["利润开始兑现", "订单增长20%"],
    )
    assert result["ok"], result

    doc = Document(path)
    h1s = [p for p in doc.paragraphs if p.style and p.style.name == "Heading 1"]
    assert len(h1s) == 1
    assert h1s[0].text.startswith("一、")
    assert all(run.font.size is not None and run.font.size.pt == 18 for p in h1s for run in p.runs)

    h2 = next(p for p in doc.paragraphs if p.text.startswith("（1）"))
    assert all(run.font.size is not None and run.font.size.pt == 15 for run in h2.runs)

    h3 = next(p for p in doc.paragraphs if p.text.startswith("①"))
    assert all(run.font.size is not None and run.font.size.pt == 12 for run in h3.runs)

    body = next(p for p in doc.paragraphs if p.text.startswith("国电南瑞在这里再次出现"))
    assert all(run.font.size is not None and run.font.size.pt == 13.5 for run in body.runs)

    risk = next(p for p in doc.paragraphs if p.text.startswith("风险提示："))
    assert all(run.italic is True for run in risk.runs if run.text)
    assert all(run.underline is False for p in doc.paragraphs for run in p.runs if run.text)
    assert not any("资料来源" in p.text for p in doc.paragraphs)


def test_a_share_company_names_are_automatically_bolded_and_skim_anchors_remain_bold(tmp_path):
    policy = json.loads(open("config/policy.json", encoding="utf-8").read())
    article = ArticleDocument(
        title="公司加粗测试",
        stock_names=["国电南瑞", "许继电气"],
        blocks=[
            Block(kind="lead", text="国电南瑞和许继电气都出现，海外公司Siemens不应被自动加粗，**利润弹性**仍是显式主线锚点。"),
            Block(kind="h1", text="比较维度"),
            Block(kind="body", text="国电南瑞再次出现，许继电气也再次出现，只有**订单兑现**和**现金流改善**属于额外扫读路径。"),
            Block(kind="h1", text="风险提示与免责声明"),
            Block(kind="body", text=_risk_text()),
        ],
    )
    path = build_article_docx(article, tmp_path / "bold.docx", policy)
    result = check_docx(
        path,
        stock_names=article.stock_names,
        policy=policy,
        required_bold_anchors=["利润弹性", "订单兑现", "现金流改善"],
    )
    assert result["ok"], result

    doc = Document(path)
    ordinary = next(p for p in doc.paragraphs if p.text.startswith("国电南瑞和许继电气"))
    assert any(run.bold and "国电南瑞" in run.text for run in ordinary.runs)
    assert any(run.bold and "许继电气" in run.text for run in ordinary.runs)
    assert any(run.bold and "利润弹性" in run.text for run in ordinary.runs)
    assert any("Siemens" in run.text and not run.bold for run in ordinary.runs)


def test_checker_rejects_any_unbold_a_share_occurrence_when_auto_bold_is_required(tmp_path):
    policy = json.loads(open("config/policy.json", encoding="utf-8").read())
    article = ArticleDocument(
        title="公司加粗机械检查",
        stock_names=["国电南瑞"],
        blocks=[
            Block(kind="lead", text="国电南瑞。"),
            Block(kind="h1", text="正文"),
            Block(kind="body", text="**关键判断**。"),
            Block(kind="h1", text="风险提示与免责声明"),
            Block(kind="body", text=_risk_text()),
        ],
    )
    path = build_article_docx(article, tmp_path / "unbold.docx", policy)
    doc = Document(path)
    p = doc.add_paragraph()
    run = p.add_run("国电南瑞")
    run.bold = False
    run.font.size = doc.styles["Normal"].font.size
    doc.save(path)

    result = check_docx(path, stock_names=article.stock_names, policy=policy, required_bold_anchors=["关键判断"])
    assert not result["ok"]
    assert any("must be bold at every occurrence" in error for error in result["errors"])


def test_concise_nonempty_risk_disclaimer_is_allowed_without_fixed_character_floor(tmp_path):
    policy = json.loads(open("config/policy.json", encoding="utf-8").read())
    assert policy["risk_disclaimer"]["min_body_chars"] == 0
    article = ArticleDocument(
        title="简洁风险提示测试",
        blocks=[
            Block(kind="lead", text="**关键事实**。"),
            Block(kind="h1", text="正文"),
            Block(kind="body", text="正文**关键判断**。"),
            Block(kind="h1", text="风险提示与免责声明"),
            Block(kind="body", text="需求与业绩可能变化，本文不构成投资建议。"),
        ],
    )
    path = build_article_docx(article, tmp_path / "risk_short.docx", policy)
    result = check_docx(path, stock_names=[], policy=policy, required_bold_anchors=["关键事实", "关键判断"])
    assert result["ok"], result


def test_tables_are_hard_fail(tmp_path):
    policy = json.loads(open("config/policy.json", encoding="utf-8").read())
    article = ArticleDocument(
        title="无表格测试",
        blocks=[
            Block(kind="lead", text="**关键事实**。"),
            Block(kind="h1", text="正文"),
            Block(kind="body", text="正文**关键判断**。"),
            Block(kind="h1", text="风险提示与免责声明"),
            Block(kind="body", text=_risk_text()),
        ],
    )
    path = build_article_docx(article, tmp_path / "table.docx", policy)
    doc = Document(path)
    doc.add_table(rows=1, cols=2)
    doc.save(path)
    result = check_docx(path, stock_names=[], policy=policy, required_bold_anchors=["关键事实", "关键判断"])
    assert not result["ok"]
    assert any("prohibited table" in error for error in result["errors"])


def test_source_appendix_and_underlines_are_hard_fail(tmp_path):
    policy = json.loads(open("config/policy.json", encoding="utf-8").read())
    article = ArticleDocument(
        title="来源与下划线测试",
        blocks=[
            Block(kind="lead", text="**关键事实**。"),
            Block(kind="h1", text="正文"),
            Block(kind="body", text="正文**关键判断**。"),
            Block(kind="h1", text="风险提示与免责声明"),
            Block(kind="body", text=_risk_text()),
        ],
    )
    path = build_article_docx(article, tmp_path / "source_underline.docx", policy)
    doc = Document(path)
    p = doc.add_paragraph()
    run = p.add_run("资料来源")
    run.underline = True
    run.font.size = doc.styles["Normal"].font.size
    doc.save(path)

    result = check_docx(path, stock_names=[], policy=policy, required_bold_anchors=["关键事实", "关键判断"])
    assert not result["ok"]
    assert any("end source section" in error for error in result["errors"])
    assert any("underline" in error for error in result["errors"])
